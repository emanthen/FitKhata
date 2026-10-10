from __future__ import annotations

import csv
import io
from decimal import Decimal
from types import SimpleNamespace

from openpyxl import load_workbook

from reports.exporters import (
    DUES_AGEING_LABELS,
    render_vat_register_pdf,
    rows_to_csv,
    rows_to_xlsx,
)

ROWS = [
    {"bucket": "0-30", "count": 2, "total_npr": Decimal("100.00")},
    {"bucket": "31-60", "count": 1, "total_npr": Decimal("200.00")},
]
FIELDS = ["bucket", "count", "total_npr"]


def test_rows_to_csv_row_count_matches_input():
    out = rows_to_csv(ROWS, FIELDS)
    reader = csv.DictReader(io.StringIO(out.decode("utf-8")))
    assert len(list(reader)) == len(ROWS)


def test_rows_to_xlsx_row_count_matches_input():
    out = rows_to_xlsx(ROWS, FIELDS, DUES_AGEING_LABELS["en"])
    ws = load_workbook(io.BytesIO(out)).active
    assert ws.max_row - 1 == len(ROWS)  # minus the header row
    assert [c.value for c in ws[1]] == [DUES_AGEING_LABELS["en"][f] for f in FIELDS]


def test_render_vat_register_pdf_returns_nonempty_bytes():
    gym = SimpleNamespace(name="Reports Test Gym")
    register = {
        "rows": [
            {
                "bill_date_bs": "2083-06-01",
                "number": "2083-00001",
                "buyer_name": "Test Member",
                "buyer_pan": "",
                "taxable_npr": Decimal("1000.00"),
                "vat_npr": Decimal("130.00"),
                "total_npr": Decimal("1130.00"),
            }
        ]
    }
    pdf_bytes = render_vat_register_pdf(gym, 2083, register)
    assert pdf_bytes.startswith(b"%PDF")
